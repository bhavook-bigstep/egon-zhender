"""NativeExtractor — real born-digital extraction (in-venv, no container).

Routes by content type: PDF → pypdf, DOCX → python-docx, XLSX → openpyxl, text → decode.
Image types defer to the OCR gate. A born-digital file with no recoverable text layer
(e.g. a scanned PDF) returns "no text layer" so the OCR gate engages. Parser failures
become a typed `EXTRACTION_ERROR` (never a silent drop). Extracted text/locations are
in-memory working data — never logged (privacy rules).
"""

from __future__ import annotations

import io

from libs.schemas import EvidenceLocation, ExceptionCode, ManifestEntry
from pipelines.workstream1_sensitive.errors import PipelineItemError
from pipelines.workstream1_sensitive.extract import (
    ExtractedText,
    decode_text,
    no_text_layer_result,
    text_result,
)
from pipelines.workstream1_sensitive.ingest import requires_ocr_type

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _assemble(source_id: str, spans: list[tuple[EvidenceLocation, str]]) -> ExtractedText:
    if not any(text.strip() for _, text in spans):
        return no_text_layer_result(source_id)  # no recoverable text → OCR gate
    text = "\n".join(text for _, text in spans)
    return text_result(source_id, text, spans, coverage_complete=True)


class NativeExtractor:
    def extract(self, entry: ManifestEntry, data: bytes) -> ExtractedText:
        if requires_ocr_type(entry):
            return no_text_layer_result(entry.source_id)
        content_type = entry.content_type
        if content_type == PDF:
            return self._pdf(entry.source_id, data)
        if content_type == DOCX:
            return self._docx(entry.source_id, data)
        if content_type == XLSX:
            return self._xlsx(entry.source_id, data)
        return decode_text(entry.source_id, data)

    def _pdf(self, source_id: str, data: bytes) -> ExtractedText:
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            spans: list[tuple[EvidenceLocation, str]] = []
            for page_index, page in enumerate(reader.pages):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    spans.append((EvidenceLocation(page=page_index + 1), page_text))
        except PipelineItemError:
            raise
        except Exception as exc:  # parser failure → typed, fail loud
            raise PipelineItemError(
                ExceptionCode.EXTRACTION_ERROR, f"pdf parse failed for {source_id}"
            ) from exc
        return _assemble(source_id, spans)

    def _docx(self, source_id: str, data: bytes) -> ExtractedText:
        from docx import Document

        try:
            document = Document(io.BytesIO(data))
            spans: list[tuple[EvidenceLocation, str]] = []
            offset = 0
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if text:
                    spans.append(
                        (
                            EvidenceLocation(char_start=offset, char_end=offset + len(text)),
                            text,
                        )
                    )
                    offset += len(text) + 1
            for table_index, table in enumerate(document.tables):
                for row_index, row in enumerate(table.rows):
                    for col_index, cell in enumerate(row.cells):
                        cell_text = cell.text.strip()
                        if cell_text:
                            spans.append(
                                (
                                    EvidenceLocation(
                                        sheet=f"table{table_index}",
                                        cell=f"r{row_index}c{col_index}",
                                    ),
                                    cell_text,
                                )
                            )
        except PipelineItemError:
            raise
        except Exception as exc:
            raise PipelineItemError(
                ExceptionCode.EXTRACTION_ERROR, f"docx parse failed for {source_id}"
            ) from exc
        return _assemble(source_id, spans)

    def _xlsx(self, source_id: str, data: bytes) -> ExtractedText:
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            spans: list[tuple[EvidenceLocation, str]] = []
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        value = str(cell.value).strip()
                        if value:
                            spans.append(
                                (
                                    EvidenceLocation(
                                        sheet=worksheet.title, cell=cell.coordinate
                                    ),
                                    value,
                                )
                            )
            workbook.close()
        except PipelineItemError:
            raise
        except Exception as exc:
            raise PipelineItemError(
                ExceptionCode.EXTRACTION_ERROR, f"xlsx parse failed for {source_id}"
            ) from exc
        return _assemble(source_id, spans)
